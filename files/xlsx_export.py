"""
Exporta as amostras gravadas para .xlsx com:
  - aba "Dados": série temporal bruta (tempo, FC, potência, cadência, etapa)
  - aba "Gráficos": 3 gráficos de linha (FC x tempo, Potência x tempo, Cadência x tempo)

Dados brutos de sensor, não um modelo financeiro — valores são gravados
diretamente (não há fórmula a preservar), então não se aplicam as
convenções de cor de modelo financeiro.
"""
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import LineChart, Reference

FONT_NAME = "Arial"


def export_xlsx(samples: list[dict], path: str):
    if not samples:
        raise ValueError("Nenhuma amostra gravada — treino não foi iniciado ou terminou sem dados.")

    wb = Workbook()
    ws_data = wb.active
    ws_data.title = "Dados"

    headers = ["Tempo (s)", "FC (bpm)", "Potência (W)", "Cadência (rpm)", "Etapa"]
    for col, h in enumerate(headers, start=1):
        cell = ws_data.cell(row=1, column=col, value=h)
        cell.font = Font(name=FONT_NAME, bold=True)

    for row_idx, s in enumerate(samples, start=2):
        ws_data.cell(row=row_idx, column=1, value=round(s["elapsed_seconds"], 1))
        ws_data.cell(row=row_idx, column=2, value=s["heart_rate"] or None)
        ws_data.cell(row=row_idx, column=3, value=s["power"] or None)
        ws_data.cell(row=row_idx, column=4, value=s["cadence"] or None)
        ws_data.cell(row=row_idx, column=5, value=s["step_name"] or "")
        for col in range(1, 6):
            ws_data.cell(row=row_idx, column=col).font = Font(name=FONT_NAME)

    for col, width in zip("ABCDE", [12, 12, 14, 14, 24]):
        ws_data.column_dimensions[col].width = width

    n_rows = len(samples)
    last_row = n_rows + 1  # +1 pelo header

    ws_charts = wb.create_sheet("Gráficos")

    time_ref = Reference(ws_data, min_col=1, min_row=2, max_row=last_row)

    def make_chart(title: str, data_col: int, y_title: str, anchor: str, color: str):
        chart = LineChart()
        chart.title = title
        chart.style = 2
        chart.x_axis.title = "Tempo (s)"
        chart.y_axis.title = y_title
        chart.height = 8
        chart.width = 22

        data_ref = Reference(ws_data, min_col=data_col, min_row=1, max_row=last_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(time_ref)

        series = chart.series[0]
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 15000  # EMU, ~1.2pt
        series.smooth = False

        ws_charts.add_chart(chart, anchor)

    make_chart("Frequência Cardíaca x Tempo", 2, "FC (bpm)", "A1", "FF4D4D")
    make_chart("Potência x Tempo", 3, "Potência (W)", "A18", "FFCC00")
    make_chart("Cadência x Tempo", 4, "Cadência (rpm)", "A35", "4DD2FF")

    wb.save(path)
